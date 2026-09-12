import logging
import io
import json

from flask import Flask, jsonify, request, send_file, send_from_directory
from flask_cors import CORS

from .config import FRONTEND_DIST, MAX_CONTENT_LENGTH, SUPPORTED_TYPES
from .database import init_db, latest_result, list_results
from .services.documents import process_document


def create_app():
    app = Flask(__name__, static_folder=str(FRONTEND_DIST) if FRONTEND_DIST.exists() else None)
    app.config["MAX_CONTENT_LENGTH"] = MAX_CONTENT_LENGTH
    CORS(app)
    logging.basicConfig(level=logging.INFO)
    init_db()

    @app.get("/api/v1/health")
    def health():
        return jsonify({"status": "ok", "service": "ledgerlens-api"})

    @app.get("/api/v1/docs")
    def docs():
        return jsonify({"openapi": "3.0.0", "info": {"title": "LedgerLens API", "version": "1.0.0"}, "paths": {"/api/v1/documents/process": {"post": {"summary": "Process a financial document"}}, "/api/v1/documents": {"get": {"summary": "List processed documents"}}}})

    @app.get("/api/v1/documents")
    def documents():
        return jsonify({"documents": list_results()})

    @app.get("/api/v1/documents/<path:document_name>")
    def document(document_name):
        result = latest_result(document_name)
        return jsonify(result or {"error": {"code": "DOCUMENT_NOT_FOUND", "message": "No processed document with that name."}}), (200 if result else 404)

    @app.get("/api/v1/documents/<path:document_name>/report")
    def document_report(document_name):
        result = latest_result(document_name)
        if not result:
            return jsonify({"error": {"code": "DOCUMENT_NOT_FOUND", "message": "No processed document with that name."}}), 404
        try:
            from docx import Document
            from docx.shared import Inches, Pt

            report = Document()
            report.add_heading("LedgerLens financial document report", 0)
            report.add_paragraph(result["document_name"])
            report.add_heading("Processing summary", level=1)
            report.add_paragraph(f"Document type: {result['document_type']}")
            report.add_paragraph(f"Processing status: {result['processing_status']}")
            report.add_paragraph(f"Validation status: {result['validation']['overall_status']}")
            report.add_heading("Financial validation analysis", level=1)
            checks = result.get("validation", {}).get("checks", [])
            if not checks:
                report.add_paragraph("No validation could be performed because the required financial values were not extracted.")
            for check in checks:
                report.add_paragraph(check.get("name", "validation check"), style="Heading 2")
                report.add_paragraph(f"Status: {check.get('status')} | Formula: {check.get('formula')}")
                report.add_paragraph(f"Calculated value: {check.get('calculated_value')} | Reported value: {check.get('reported_value')} | Variance: {check.get('variance')}")
                report.add_paragraph(f"Operands: {json.dumps(check.get('operands', {}), ensure_ascii=False)}")
            report.add_heading("Extracted data", level=1)
            for field, item in result.get("extracted_data", {}).items():
                value = item.get("value") if isinstance(item, dict) else item
                evidence = item.get("evidence") if isinstance(item, dict) else None
                paragraph = report.add_paragraph(style="List Bullet")
                paragraph.add_run(f"{field}: {value}").bold = True
                if evidence:
                    paragraph.add_run(f" | Evidence: {evidence}")
            report.add_heading("Structured JSON", level=1)
            json_paragraph = report.add_paragraph()
            json_run = json_paragraph.add_run(json.dumps(result, indent=2, ensure_ascii=False))
            json_run.font.name = "Consolas"
            json_run.font.size = Pt(8)
            output = io.BytesIO()
            report.save(output)
            output.seek(0)
            return send_file(output, as_attachment=True, download_name=f"{document_name.rsplit('.', 1)[0]}_report.docx", mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        except Exception:
            app.logger.exception("Report generation failed")
            return jsonify({"error": {"code": "REPORT_ERROR", "message": "The Word report could not be generated."}}), 500

    @app.post("/api/v1/documents/process")
    def process():
        uploaded = request.files.get("file")
        document_type = request.form.get("document_type", "")
        if not uploaded or not uploaded.filename:
            return jsonify({"error": {"code": "FILE_REQUIRED", "message": "Attach a PDF, JPG, JPEG, or PNG file."}}), 400
        if document_type not in SUPPORTED_TYPES:
            return jsonify({"error": {"code": "INVALID_DOCUMENT_TYPE", "message": "Choose one of the supported financial document types."}}), 400
        try:
            return jsonify(process_document(uploaded, document_type)), 201
        except ValueError as error:
            return jsonify({"error": {"code": "INVALID_DOCUMENT", "message": str(error)}}), 400
        except Exception:
            app.logger.exception("Document processing failed")
            return jsonify({"error": {"code": "PROCESSING_ERROR", "message": "The document could not be processed. Check OCR and model configuration."}}), 500

    @app.errorhandler(413)
    def too_large(_error):
        return jsonify({"error": {"code": "FILE_TOO_LARGE", "message": "The upload exceeds the 15 MB limit."}}), 413

    if FRONTEND_DIST.exists():
        @app.get("/")
        def frontend_index():
            return send_from_directory(FRONTEND_DIST, "index.html")

        @app.get("/<path:asset>")
        def frontend_assets(asset):
            return send_from_directory(FRONTEND_DIST, asset)
    return app


app = create_app()
